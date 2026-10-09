"""Local Office layout/raster extraction contract for artifact renderers."""
import base64
import io
from pathlib import Path
import random
import sys
import shutil
import uuid
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import office_preview as preview
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from docx import Document
from docx.shared import Inches as DocInches


def raster(fmt='PNG', size=(2200, 120)):
    stream = io.BytesIO()
    Image.new('RGB', size, (140, 30, 20)).save(stream, fmt)
    return stream.getvalue()


class ArtifactPreviewTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(__file__).resolve().parents[1] / 'data' / ('artifact-preview-' + uuid.uuid4().hex)
        self.root.mkdir()
        self.addCleanup(shutil.rmtree, self.root)
        self.picture = self.root / 'picture.png'
        self.picture.write_bytes(raster())

    def deck(self):
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(12), Inches(6.75)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = 'Bold red'; run.font.bold = True; run.font.italic = True
        run.font.underline = True; run.font.size = Pt(24); run.font.color.rgb = RGBColor(255, 0, 0)
        shape.rotation = 30
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(10, 20, 30)
        shape.line.color.rgb = RGBColor(1, 2, 3)
        picture = slide.shapes.add_picture(str(self.picture), Inches(5), Inches(1), width=Inches(2))
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(3), Inches(1)).table
        table.cell(0, 0).text = 'Cell A'; table.cell(1, 1).text = 'Cell B'
        group = slide.shapes.add_group_shape()
        one = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        two = group.shapes.add_textbox(Inches(3), Inches(1), Inches(1), Inches(1))
        two.text = 'Grouped'
        group.left, group.top, group.width, group.height = Inches(2), Inches(3), Inches(6), Inches(2)
        theme = slide.shapes.add_textbox(Inches(8), Inches(2), Inches(2), Inches(1))
        theme.text = 'Theme'; theme.text_frame.paragraphs[0].runs[0].font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        theme.fill.solid(); theme.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(240, 230, 220)
        slide.notes_slide.notes_text_frame.text = 'Speaker notes'
        for label in ('Slide two', 'Slide three'):
            other = deck.slides.add_slide(deck.slide_layouts[6])
            other.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = label
        path = self.root / 'deck.pptx'; deck.save(path)
        return path, shape.shape_id, picture.shape_id, group.shape_id, one.shape_id, two.shape_id, theme.shape_id

    def test_deck_geometry_runs_groups_and_legacy_fields(self):
        path, text_id, picture_id, group_id, first, second, theme_id = self.deck()
        data = preview.extract_pptx(path)
        self.assertEqual((data['slide_width'], data['slide_height']), (Inches(12), Inches(6.75)))
        self.assertEqual(len(data['slides']), 3)
        slide = data['slides'][0]
        self.assertTrue({'title', 'body', 'tables', 'notes'}.issubset(slide))
        self.assertIn('Bold red', slide['body']); self.assertIn('Speaker notes', slide['notes'])
        self.assertEqual(slide['tables'][0][0][0], 'Cell A')
        self.assertEqual(slide['background'], '#f0e6dc')
        shapes = {shape['id']: shape for shape in slide['shapes']}
        text = shapes[text_id]
        self.assertEqual((text['x'], text['y'], text['w'], text['h']), (Inches(1), Inches(2), Inches(3), Inches(1)))
        self.assertEqual(text['rotation'], 30)
        self.assertEqual(text['fill'], '#0a141e'); self.assertEqual(text['line'], '#010203')
        run = text['paragraphs'][0]['runs'][0]
        self.assertEqual(run, {'text': 'Bold red', 'size_pt': 24, 'bold': True, 'italic': True, 'underline': True, 'color': '#ff0000'})
        self.assertNotIn(group_id, shapes)
        for child in (first, second): self.assertEqual(shapes[child]['group'], group_id)
        self.assertEqual((shapes[first]['x'], shapes[first]['y'], shapes[first]['w'], shapes[first]['h']), (Inches(2), Inches(3), Inches(2), Inches(2)))
        self.assertEqual(shapes[second]['x'], Inches(6))
        self.assertEqual([s['z'] for s in slide['shapes']], list(range(len(shapes))))
        self.assertIsNone(shapes[theme_id]['paragraphs'][0]['runs'][0]['color'])
        self.assertIsNone(shapes[theme_id]['fill'])
        uri = shapes[picture_id]['image']; self.assertTrue(uri.startswith('data:image/png;base64,'))
        with Image.open(io.BytesIO(base64.b64decode(uri.split(',')[1]))) as img:
            self.assertLessEqual(max(img.size), 1600)

    def test_docx_formatted_lists_and_inline_body_order(self):
        document = Document(); document.add_heading('Heading', 1)
        paragraph = document.add_paragraph(); run = paragraph.add_run('Bold'); run.bold = True; run.italic = True; run.underline = True
        document.add_paragraph('Bullet', 'List Bullet'); document.add_paragraph('Number', 'List Number')
        paragraph = document.add_paragraph('Before '); paragraph.add_run().add_picture(str(self.picture), width=DocInches(2)); paragraph.add_run(' After')
        document.add_table(1, 1).cell(0, 0).text = 'Table'
        path = self.root / 'doc.docx'; document.save(path)
        blocks = preview.extract_docx(path)['blocks']
        self.assertEqual([b['type'] for b in blocks], ['heading', 'paragraph', 'list', 'list', 'paragraph', 'image', 'paragraph', 'table'])
        self.assertTrue(all('text' in block for block in blocks))
        self.assertEqual(blocks[1]['runs'][0], {'text': 'Bold', 'bold': True, 'italic': True, 'underline': True})
        self.assertEqual(blocks[2]['list'], 'bullet'); self.assertEqual(blocks[3]['list'], 'number')
        self.assertEqual(blocks[2]['level'], 0)
        self.assertTrue(blocks[5]['image'].startswith('data:image/png;base64,'))
        self.assertEqual(blocks[4]['text'], 'Before '); self.assertEqual(blocks[6]['text'], ' After')

    def replace_picture(self, source, suffix):
        path = self.root / ('hostile' + suffix)
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, 'w') as updated:
            for entry in original.infolist():
                updated.writestr(entry, b'\x01\x00\x00\x00 EMF not a raster' if '/media/' in entry.filename else original.read(entry))
        return path

    def test_non_raster_blob_becomes_picture_placeholder_in_both_formats(self):
        path, _, picture_id, *_ = self.deck()
        data = preview.extract_pptx(self.replace_picture(path, '.pptx'))
        shape = next(s for s in data['slides'][0]['shapes'] if s['id'] == picture_id)
        self.assertEqual(shape['kind'], 'picture'); self.assertIsNone(shape['image'])
        document = Document(); document.add_picture(str(self.picture)); path = self.root / 'doc.docx'; document.save(path)
        block = preview.extract_docx(self.replace_picture(path, '.docx'))['blocks'][0]
        self.assertEqual(block['type'], 'image'); self.assertIsNone(block['image']); self.assertIn('unavailable', block['text'])

    def test_magic_formats_reencoding_and_caps(self):
        for fmt in ('PNG', 'JPEG', 'GIF', 'WEBP'):
            self.assertIsNotNone(preview._ImageBudget().read(raster(fmt, (20, 20))), fmt)
        for blob in (b'<svg/>', b'EMF', b'GIF89afake', b'\x89PNG\r\n\x1a\nfake'):
            self.assertIsNone(preview._ImageBudget().read(blob))
        blob = raster('PNG', (20, 20)); uri = preview._ImageBudget().read(blob); size = len(uri)
        with patch.object(preview, 'MAX_IMAGE_BYTES', size - 1):
            self.assertIsNone(preview._ImageBudget().read(blob))
        with patch.object(preview, 'MAX_TOTAL_IMAGE_BYTES', size * 2):
            budget = preview._ImageBudget()
            self.assertIsNotNone(budget.read(blob)); self.assertIsNotNone(budget.read(blob)); self.assertIsNone(budget.read(blob)); self.assertEqual(budget.used, size * 2)
        # Real encoded payload above the configured ceiling, with no mocked encoder.
        # An opaque photo saved as PNG steps down to JPEG instead of vanishing.
        noise = Image.frombytes('RGB', (1800, 1800), random.Random(0).randbytes(1800 * 1800 * 3)); stream = io.BytesIO(); noise.save(stream, 'PNG')
        photo = preview._ImageBudget().read(stream.getvalue())
        self.assertTrue(photo.startswith('data:image/jpeg;base64,')); self.assertLessEqual(len(photo), preview.MAX_IMAGE_BYTES)
        # Transparent noise cannot become JPEG and stays over the cap even at 800 px.
        alpha = Image.frombytes('RGBA', (900, 900), random.Random(1).randbytes(900 * 900 * 4)); stream = io.BytesIO(); alpha.save(stream, 'PNG')
        self.assertIsNone(preview._ImageBudget().read(stream.getvalue()))
        path, *_ = self.deck()
        with patch.object(preview, 'MAX_TOTAL_IMAGE_BYTES', 1):
            self.assertTrue(all(s['image'] is None for slide in preview.extract_pptx(path)['slides'] for s in slide['shapes']))

    def test_archive_and_walk_caps_still_apply(self):
        path = self.root / 'bomb.pptx'
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive: archive.writestr('bomb', b'0' * 1000000)
        for extract in (preview.extract_pptx, preview.extract_docx, preview.extract_xlsx):
            with self.assertRaises(preview.PreviewUnavailable): extract(path)
        path, *_ = self.deck()
        with patch.object(preview, 'MAX_SLIDES', 1), patch.object(preview, 'MAX_SHAPES_PER_SLIDE', 2):
            data = preview.extract_pptx(path); self.assertEqual(len(data['slides']), 1); self.assertEqual(len(data['slides'][0]['shapes']), 2); self.assertTrue(data['truncated'])
        with patch.object(preview, 'MAX_CELL_CHARS', 4):
            data = preview.extract_pptx(path)
            self.assertEqual(data['slides'][0]['shapes'][0]['paragraphs'][0]['runs'][0]['text'], 'Bold')

    def test_image_budget_spans_slides_and_docx_blocks(self):
        path, *_ = self.deck()
        deck = Presentation(path)
        for slide in deck.slides:
            slide.shapes.add_picture(str(self.picture), Inches(1), Inches(1), width=Inches(1))
        deck.save(path)
        document = Document()
        for _ in range(3): document.add_picture(str(self.picture))
        docpath = self.root / 'pictures.docx'; document.save(docpath)
        size = len(preview._ImageBudget().read(self.picture.read_bytes()))
        with patch.object(preview, 'MAX_TOTAL_IMAGE_BYTES', size * 2):
            slides = preview.extract_pptx(path)['slides']
            pictures = [s for slide in slides for s in slide['shapes'] if s['kind'] == 'picture']
            self.assertEqual(sum(s['image'] is not None for s in pictures), 2)
            self.assertLessEqual(sum(len(s['image'] or '') for s in pictures), size * 2)
            blocks = preview.extract_docx(docpath)['blocks']
            self.assertEqual(sum(b['image'] is not None for b in blocks), 2)
        with patch.object(preview, 'MAX_IMAGE_BYTES', size - 1):
            self.assertTrue(all(s['image'] is None for slide in preview.extract_pptx(path)['slides'] for s in slide['shapes']))
            self.assertTrue(all(b['image'] is None for b in preview.extract_docx(docpath)['blocks']))

    def test_picture_placeholder_carries_its_image(self):
        deck = Presentation()
        layout = next(l for l in deck.slide_layouts if any(ph.placeholder_format.type == 18 for ph in l.placeholders))
        slide = deck.slides.add_slide(layout)
        holder = next(ph for ph in slide.placeholders if ph.placeholder_format.type == 18)
        filled = holder.insert_picture(str(self.picture))
        path = self.root / 'photo.pptx'; deck.save(path)
        shape = next(s for s in preview.extract_pptx(path)['slides'][0]['shapes'] if s['id'] == filled.shape_id)
        self.assertEqual(shape['kind'], 'picture'); self.assertTrue(shape['image'].startswith('data:image/'))

    def test_rotated_group_absolute_centers(self):
        path, _, _, _, first, second, _ = self.deck()
        deck = Presentation(path)
        group = next(shape for shape in deck.slides[0].shapes if shape.shape_type == 6)
        group.rotation = 90; deck.save(path)
        shapes = {s['id']: s for s in preview.extract_pptx(path)['slides'][0]['shapes']}
        self.assertAlmostEqual(shapes[first]['x'], Inches(4), delta=1)
        self.assertAlmostEqual(shapes[first]['y'], Inches(1), delta=1)
        self.assertAlmostEqual(shapes[second]['x'], Inches(4), delta=1)
        self.assertAlmostEqual(shapes[second]['y'], Inches(5), delta=1)
        self.assertEqual(shapes[first]['rotation'], 90)


if __name__ == '__main__': unittest.main()
